# -*- coding: utf-8 -*-
"""本地声优库的离线测试：store 的格式与读写，以及查看器的四种视角。

零联网、零真实用户状态：库一律建在 tmp_path 下。
"""
import dataclasses
import json
import os
import threading

import pytest

from galtools.core.context import RunContext
from galtools.tools import seiyuu_db as tool
from galtools.tools.vndb_voiced import store, xlsx
from galtools.tools.vndb_voiced.model import Credit, Staff, StaffCredits


# ---------------- 造数据 ----------------
def credit(vid='v1', **kw):
    base = dict(title='Title %s' % vid, title_ja='原題 %s' % vid,
                released='2012-08-31', cid='c9', cast='Cast', cast_ja='キャスト',
                alias='Alias', alias_ja='別名', note='note', role='main')
    base.update(kw)
    return Credit(vid=vid, **base)


def person(sid='s1', name='Alpha One', original='あるふぁ', credits=None):
    return StaffCredits(
        staff=Staff(sid=sid, name=name, original=original),
        credits=[credit()] if credits is None else credits)


def dump(db_dir, name, payload):
    """直接落一个文件，用来造工具自己写不出来的形状。"""
    path = os.path.join(str(db_dir), name)
    text = payload if isinstance(payload, str) else json.dumps(
        payload, ensure_ascii=False)
    with open(path, 'w', encoding='utf-8') as fp:
        fp.write(text)
    return path


# ---------------- 格式 ----------------
def test_credit_fields_cover_the_dataclass():
    """给 Credit 加字段就得顺手决定库格式跟不跟——这条断言是那个提醒。"""
    assert store.CREDIT_FIELDS == tuple(
        f.name for f in dataclasses.fields(Credit))


def test_round_trip_keeps_every_field(tmp_path):
    item = person(credits=[credit('v1'), credit('v2', role='side', note='')])
    store.write_person(str(tmp_path), item, fetched_at='2026-09-04T20:11:33')

    back = store.read_person(str(tmp_path), 's1')
    assert back.fetched_at == '2026-09-04T20:11:33'
    assert back.staff == Staff(sid='s1', name='Alpha One', original='あるふぁ')
    assert back.credits == item.credits


def test_written_file_is_readable_by_a_human(tmp_path):
    store.write_person(str(tmp_path), person())
    text = (tmp_path / 's1.json').read_text(encoding='utf-8')
    # 缩进两格 + 不转义非 ASCII：这个库是用户会自己打开看的。
    assert '\n  "sid": "s1"' in text
    assert 'あるふぁ' in text
    assert json.loads(text)['schema'] == store.SCHEMA


def test_credits_keep_their_order(tmp_path):
    vids = ['v3', 'v1', 'v2']
    store.write_person(str(tmp_path), person(credits=[credit(v) for v in vids]))
    assert [c.vid for c in store.read_person(str(tmp_path), 's1').credits] == vids


def test_aliases_are_not_stored(tmp_path):
    item = person()
    item.staff = Staff(sid='s1', name='Alpha One', original='あるふぁ',
                       aliases={'a1': ('Alias', '別名')})
    store.write_person(str(tmp_path), item)
    data = json.loads((tmp_path / 's1.json').read_text(encoding='utf-8'))
    assert 'aliases' not in data


def test_second_write_replaces_in_place(tmp_path):
    store.write_person(str(tmp_path), person(credits=[credit('v1')]))
    store.write_person(str(tmp_path), person(credits=[credit('v2')]))
    assert sorted(os.listdir(str(tmp_path))) == ['s1.json']
    assert [c.vid for c in store.read_person(str(tmp_path), 's1').credits] == ['v2']


def test_sid_must_look_like_a_vndb_id(tmp_path):
    """文件名要拼进路径，sid 是唯一的来源。"""
    for bad in ('', '../s1', 'x367', 's367 '):
        item = person(sid=bad)
        with pytest.raises(store.BadFile):
            store.write_person(str(tmp_path), item)
    assert os.listdir(str(tmp_path)) == []


def test_failed_write_leaves_no_tmp_behind(tmp_path, monkeypatch):
    store.write_person(str(tmp_path), person(credits=[credit('v1')]))

    def boom(*_args, **_kw):
        raise OSError('盘满了')

    monkeypatch.setattr(store.json, 'dump', boom)
    with pytest.raises(OSError):
        store.write_person(str(tmp_path), person(credits=[credit('v2')]))
    # 上一份完好，目录里没有多出来的碎片。
    assert sorted(os.listdir(str(tmp_path))) == ['s1.json']
    assert [c.vid for c in store.read_person(str(tmp_path), 's1').credits] == ['v1']


def test_concurrent_writers_use_separate_temporary_files(tmp_path, monkeypatch):
    real_dump = store.json.dump
    real_replace = store.os.replace
    barrier = threading.Barrier(2)
    release_first = threading.Event()
    replaced = []

    def together(obj, fp, **kw):
        barrier.wait(timeout=5)
        return real_dump(obj, fp, **kw)

    def serialized_replace(src, dst):
        if not replaced:
            replaced.append(src)
            result = real_replace(src, dst)
            release_first.set()
            return result
        assert release_first.wait(5)
        return real_replace(src, dst)

    monkeypatch.setattr(store.json, 'dump', together)
    monkeypatch.setattr(store.os, 'replace', serialized_replace)
    errors = []

    def write(vid):
        try:
            store.write_person(str(tmp_path), person(credits=[credit(vid)]))
        except Exception as e:
            errors.append(e)

    threads = [threading.Thread(target=write, args=(vid,))
               for vid in ('v1', 'v2')]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(5)

    assert not errors
    assert all(not thread.is_alive() for thread in threads)
    assert sorted(os.listdir(str(tmp_path))) == ['s1.json']
    assert store.read_person(str(tmp_path), 's1').credits[0].vid in ('v1', 'v2')


# ---------------- 容错 ----------------
def test_unknown_keys_are_ignored(tmp_path):
    dump(tmp_path, 's1.json', {'schema': 1, 'sid': 's1', 'name': 'Alpha One',
                               'credits': [{'vid': 'v1', 'title': 'T',
                                            '将来的字段': 42}],
                               '顶层的将来字段': ['x']})
    got = store.read_person(str(tmp_path), 's1')
    assert got.staff.sid == 's1'
    assert [c.vid for c in got.credits] == ['v1']


def test_file_name_is_the_only_identity_source(tmp_path):
    dump(tmp_path, 's1.json', {'schema': 1, 'sid': 's2', 'name': 'Wrong'})
    with pytest.raises(store.BadFile, match='文件名.*s1.*内容.*s2'):
        store.read_person(str(tmp_path), 's1')

    people, failures = store.read_all(str(tmp_path))
    assert people == []
    assert [name for name, _ in failures] == ['s1.json']


def test_schema_from_the_future_is_refused(tmp_path):
    dump(tmp_path, 's1.json', {'schema': store.SCHEMA + 1, 'sid': 's1'})
    with pytest.raises(store.BadFile) as e:
        store.read_person(str(tmp_path), 's1')
    assert '比工具新' in str(e.value)


def test_non_string_values_fall_back_to_defaults(tmp_path):
    """手改过的文件不该让读者崩在 None.casefold() 上。"""
    dump(tmp_path, 's1.json', {'sid': 's1', 'name': None, 'original': 12,
                               'credits': [{'vid': 'v1', 'title': None,
                                            'role': 'main'},
                                           'これは dict ではない']})
    got = store.read_person(str(tmp_path), 's1')
    assert (got.staff.name, got.staff.original) == ('', '')
    assert [(c.vid, c.title, c.role) for c in got.credits] == [('v1', '', 'main')]


def test_broken_file_only_blames_itself(tmp_path):
    for sid in ('s1', 's2'):
        store.write_person(str(tmp_path), person(sid=sid, name='Name ' + sid))
    dump(tmp_path, 's3.json', '{ 这不是 json')
    dump(tmp_path, 's4.json', ['顶层是个数组'])

    people, failures = store.read_all(str(tmp_path))
    assert [p.staff.sid for p in people] == ['s1', 's2']
    assert [name for name, _ in failures] == ['s3.json', 's4.json']
    assert 'JSON' in dict(failures)['s3.json']
    assert '对象' in dict(failures)['s4.json']


def test_only_s_json_files_count(tmp_path):
    store.write_person(str(tmp_path), person())
    dump(tmp_path, 'notes.json', {'sid': 's9'})
    dump(tmp_path, 's2.json.tmp', {'sid': 's2'})
    dump(tmp_path, '共同出演.xlsx', 'binary-ish')
    people, failures = store.read_all(str(tmp_path))
    assert [p.staff.sid for p in people] == ['s1']
    assert failures == []


def test_read_all_is_sorted_by_romaji(tmp_path):
    for sid, name in [('s10', 'Charlie'), ('s2', 'alpha'), ('s3', 'Bravo')]:
        store.write_person(str(tmp_path), person(sid=sid, name=name))
    people, _ = store.read_all(str(tmp_path))
    assert [p.staff.name for p in people] == ['alpha', 'Bravo', 'Charlie']


def test_missing_db_dir_is_a_bad_file(tmp_path):
    with pytest.raises(store.BadFile) as e:
        store.read_all(str(tmp_path / '没有这个目录'))
    assert '库目录' in str(e.value)


# ---------------- 查看器：造库 ----------------
def library(tmp_path, *items):
    """把几个人写进 tmp_path 下的库目录，返回它的路径。"""
    db = tmp_path / 'db'
    db.mkdir(exist_ok=True)
    for item in items:
        store.write_person(str(db), item, fetched_at='2026-09-01T12:00:00')
    return str(db)


def trio(tmp_path):
    """s1 与 s2 共演 v1、v2，s3 只在 v1。两个人姓 Ono——歧义要靠它。"""
    return library(
        tmp_path,
        person('s1', 'Ono Ryouko', '小野 涼子', [
            credit('v1', title='Alpha', title_ja='アルファ', cid='c1',
                   cast='Aiko', cast_ja='愛子'),
            credit('v2', title='Beta', title_ja='ベータ', cid='c2',
                   cast='Bell', cast_ja='ベル', role='side'),
            credit('v3', title='Gamma', title_ja='ガンマ', cid='c3',
                   cast='Cana', cast_ja='カナ', role='')]),
        person('s2', 'Mizuhashi Kaori', '水橋 かおり', [
            credit('v1', title='Alpha', title_ja='アルファ', cid='c4',
                   cast='Dora', cast_ja='ドラ'),
            credit('v2', title='Beta', title_ja='ベータ', cid='c5',
                   cast='Eri', cast_ja='エリ', role='primary')]),
        person('s3', 'Ono Kanako', '小野 佳那子', [
            credit('v1', title='Alpha', title_ja='アルファ', cid='c6',
                   cast='Fumi', cast_ja='フミ', role='appears')]))


def args(db_dir, **kw):
    """一份完整的参数字典。默认只看不导：绝大多数用例验的是屏幕上那一份。"""
    params = {'db_dir': db_dir, 'who': '', 'keyword': '', 'out_dir': ''}
    params.update(kw)
    return params


def look(db_dir, **kw):
    return tool.preview(args(db_dir, **kw), RunContext())


# ---------------- 名册 ----------------
def test_roster_shows_everyone_with_the_overview_columns(tmp_path):
    got = look(trio(tmp_path))
    assert '库里 3 人，3 部作品，6 条出演记录' in got.summary
    assert got.table.columns == ('声优', '罗马字', 'ID', '作品数', '角色数',
                                 'main', 'primary', 'side', 'appears',
                                 '未标注', '抓取时间')
    # 按罗马字排序，与 store.read_all 一致。
    assert [row[1] for row in got.table.rows] == [
        'Mizuhashi Kaori', 'Ono Kanako', 'Ono Ryouko']
    ryouko = got.table.rows[2]
    assert ryouko[2] == ('s1', 'https://vndb.org/s1')
    assert ryouko[3:] == (3, 3, 1, 0, 1, 0, 1, '2026-09-01T12:00:00')


def test_roster_never_computes_combinations(tmp_path):
    """名册可以有 50 个人，2^50 个组合是算不出来的。"""
    view = tool.build_view(args(trio(tmp_path)))
    assert view.mode == tool.ROSTER
    assert view.groups == []


def test_a_broken_file_does_not_hide_the_rest(tmp_path):
    db = library(tmp_path, person('s1', 'Alpha One'))
    dump(db, 's9.json', '{ 这不是 json')
    got = look(db)
    assert '库里 1 人' in got.summary
    assert got.warnings == ['1 个文件读不出来，已跳过：s9.json']


def test_an_empty_directory_says_where_a_library_comes_from(tmp_path):
    db = tmp_path / 'db'
    db.mkdir()
    got = look(str(db))
    assert not got.ok
    assert 'vndb 声优出演表' in got.summary


# ---------------- 挑人 ----------------
def test_who_matches_by_sid_then_full_name_then_substring(tmp_path):
    people, _ = store.read_all(trio(tmp_path))
    for who in ('s2', 'Mizuhashi Kaori', '水橋かおり', 'mizuhashi', '水橋'):
        selected, problems = tool.find_people(people, who)
        assert problems == []
        assert [p.staff.sid for p in selected] == ['s2'], who


def test_who_refuses_a_substring_that_fits_two_people(tmp_path):
    """两个人都姓 Ono。猜错了导出的是另一个人的表，而表里看不出这件事。"""
    got = look(trio(tmp_path), who='ono')
    assert not got.ok
    assert '对上了 2 个人' in got.summary
    assert 'Ono Kanako' in got.summary and 'Ono Ryouko' in got.summary
    # 退回名册：用户接下来要看的是库里到底有谁。
    assert got.table.title == '库里共 3 人'


def test_a_name_nobody_matches_falls_back_to_the_roster(tmp_path):
    got = look(trio(tmp_path), who='s404, 誰か')
    assert not got.ok
    assert got.summary.startswith('s404：库里没有这个人')
    assert got.warnings[-1] == '2 个名字没对上库里的人，先改对再导出。'
    assert got.table.title == '库里共 3 人'


def test_one_bad_name_does_not_hide_the_good_one(tmp_path):
    got = look(trio(tmp_path), who='s1, s404')
    assert not got.ok            # 少一个人的表不能悄悄导出去
    assert 's404：库里没有这个人' in got.summary
    assert got.table.columns == ('发售日', 'Title', '角色', 'As', 'Role')
    assert len(got.table.rows) == 3


def test_the_same_person_named_twice_is_merged(tmp_path):
    view = tool.build_view(args(trio(tmp_path), who='s1, Ono Ryouko'))
    assert [i.staff.sid for i in view.items] == ['s1']


def test_one_person_gets_his_own_credits(tmp_path):
    got = look(trio(tmp_path), who='s1')
    assert '小野 涼子 (Ono Ryouko) s1 : 3 部作品 / 3 个角色' in got.summary
    assert got.table.title == '小野 涼子 (Ono Ryouko) s1：3 条出演记录'


def test_two_people_get_their_common_works_offline(tmp_path):
    got = look(trio(tmp_path), who='s1, s2')
    assert '共同出演 : 2 部' in got.summary
    assert got.table.columns == ('发售日', 'Title', '日文原名',
                                 'Ono Ryouko', 'Mizuhashi Kaori')
    assert [row[1][0] for row in got.table.rows] == ['Alpha', 'Beta']
    assert got.table.rows[0][3] == ('愛子', 'https://vndb.org/c1')


def test_three_people_list_every_combination(tmp_path):
    got = look(trio(tmp_path), who='s1, s2, s3')
    assert '共同出演 : 4 个组合有交集（共 4 个组合）' in got.summary
    assert 'Ono Ryouko、Mizuhashi Kaori : 2 部' in got.summary
    assert got.table.title == '共同出演 1 部'      # 全员那一档


# ---------------- 关键词 ----------------
def test_keyword_looks_at_both_scripts_of_every_text_field(tmp_path):
    db = library(tmp_path, person('s1', 'Alpha One', credits=[
        credit('v1', title='Deep Blue', title_ja='深い青', cast='Aiko',
               cast_ja='愛子', alias='Nanashi', alias_ja='名無し', note='ノート'),
        credit('v2', title='Other', title_ja='別の', cast='Zed', cast_ja='ゼド',
               alias='', alias_ja='', note='')]))
    for keyword in ('deep', '深い青', 'aiko', '愛子', 'nanashi', '名無し',
                    'ノート'):
        view = tool.build_view(args(db, keyword=keyword))
        assert [c.vid for i in view.items for c in i.credits] == ['v1'], keyword


def test_keyword_ignores_case_and_spacing(tmp_path):
    db = library(tmp_path, person('s1', 'Alpha One', credits=[
        credit('v1', title='Deep Blue', cast_ja='水橋 かおり')]))
    for keyword in ('DEEP blue', 'deepblue', '水橋かおり', ' 水橋 かおり '):
        view = tool.build_view(args(db, keyword=keyword))
        assert len(view.items) == 1, keyword


def test_keyword_does_not_match_the_role_or_the_date(tmp_path):
    """'main' 会命中半个库，日期用不着模糊匹配。"""
    db = library(tmp_path, person('s1', 'Alpha One',
                                  credits=[credit('v1', role='main',
                                                  released='2012-08-31')]))
    for keyword in ('main', '2012'):
        assert tool.build_view(args(db, keyword=keyword)).items == []


def test_a_keyword_alone_searches_the_whole_library(tmp_path):
    got = look(trio(tmp_path), keyword='beta')
    assert '「beta」命中 2 条，2 个人（库里共 3 人）' in got.summary
    assert got.table.columns == ('声优', '发售日', 'Title', '角色', 'As', 'Role')
    assert [row[0] for row in got.table.rows] == ['水橋 かおり', '小野 涼子']


def test_people_with_no_hits_are_dropped_only_when_nobody_was_named(tmp_path):
    db = trio(tmp_path)
    # 全库搜索：0 条的人不摆，几十行「0 条」没有信息量。
    assert [i.staff.sid
            for i in tool.build_view(args(db, keyword='gamma')).items] == ['s1']
    # 点名要看的人即使 0 条也得在，否则用户会以为是自己填错了。
    view = tool.build_view(args(db, who='s1, s2', keyword='gamma'))
    assert [(i.staff.sid, len(i.credits)) for i in view.items] == [('s1', 1),
                                                                  ('s2', 0)]


def test_a_keyword_that_hits_nothing_has_nothing_to_export(tmp_path):
    got = look(trio(tmp_path), keyword='zzz', out_dir=str(tmp_path))
    assert not got.ok
    assert '没有能导出的东西。' in got.summary
    assert got.table.rows == []


# ---------------- 导出 ----------------
def test_export_names_say_which_view_they_are(tmp_path):
    """选人导出沿用抓取工具的写法：同样几个人，离线与在线该落到同一个文件名。"""
    db = trio(tmp_path)

    def name(**kw):
        return tool.workbook_name(tool.build_view(args(db, **kw)))

    assert name() == '声优库_3人.xlsx'
    assert name(who='s1') == 'vndb_Ono_Ryouko_voiced.xlsx'
    assert name(who='s1, s2') == '共同出演_Ono_Ryouko_Mizuhashi_Kaori.xlsx'
    assert name(keyword='beta') == '声优库_搜索_beta.xlsx'
    # 关键词进文件名要过 slug：空格与标点都不能直接落到路径里。
    assert name(keyword='Deep Blue!') == '声优库_搜索_Deep_Blue.xlsx'


def test_export_writes_the_roster_without_a_combo_sheet(tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / 'out'
    out.mkdir()
    result = tool.run(args(trio(tmp_path), out_dir=str(out)), RunContext())
    assert result.output_paths == [str(out / '声优库_3人.xlsx')]

    wb = load_workbook(result.output_paths[0])
    assert wb.sheetnames == ['概览', '水橋 かおり s2', '小野 佳那子 s3',
                             '小野 涼子 s1']
    assert wb['概览']['A2'].value == '水橋 かおり'


def test_export_of_two_people_carries_the_common_sheet(tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / 'out'
    out.mkdir()
    result = tool.run(args(trio(tmp_path), who='s1, s2', out_dir=str(out)),
                      RunContext())
    wb = load_workbook(result.output_paths[0])
    assert wb.sheetnames == ['概览', '共同出演', '小野 涼子 s1', '水橋 かおり s2']
    assert wb['共同出演']['B2'].value == 'Alpha'


def test_export_does_not_overwrite_an_earlier_one(tmp_path):
    out = tmp_path / 'out'
    out.mkdir()
    db = trio(tmp_path)
    first = tool.run(args(db, out_dir=str(out)), RunContext())
    second = tool.run(args(db, out_dir=str(out)), RunContext())
    assert os.path.basename(second.output_paths[0]) == '声优库_3人_1.xlsx'
    assert os.path.exists(first.output_paths[0])


def test_run_writes_nothing_without_an_export_dir(tmp_path):
    result = tool.run(args(trio(tmp_path)), RunContext())
    assert result.output_paths == []
    assert result.failures == [('导出目录', '没填')]


def test_run_recreates_an_export_dir_that_vanished_after_preview(tmp_path):
    """预览时目录还在、执行时没了（被删或拔盘）不该挡住导出：run 自己把目录建回来。"""
    out = tmp_path / 'out' / 'deeper'
    result = tool.run(args(trio(tmp_path), out_dir=str(out)), RunContext())
    assert os.path.basename(result.output_paths[0]) == '声优库_3人.xlsx'


def test_run_reports_an_export_dir_it_cannot_create(tmp_path):
    blocker = tmp_path / 'blocker'
    blocker.write_text('x', encoding='utf-8')
    result = tool.run(args(trio(tmp_path), out_dir=str(blocker / 'out')),
                      RunContext())
    assert result.output_paths == []
    assert result.failures[0][0] == '导出目录'
    assert result.failures[0][1] != ''


def test_run_refuses_when_a_name_did_not_match(tmp_path):
    out = tmp_path / 'out'
    out.mkdir()
    result = tool.run(args(trio(tmp_path), who='s1, s404', out_dir=str(out)),
                      RunContext())
    assert result.output_paths == []
    assert result.failures == [('s404', '库里没有这个人')]
    assert os.listdir(str(out)) == []


def test_run_reports_a_write_failure_instead_of_a_traceback(tmp_path,
                                                           monkeypatch):
    out = tmp_path / 'out'
    out.mkdir()

    def boom(*_args, **_kw):
        raise OSError('盘满了')

    monkeypatch.setattr(tool.xlsx, 'build', boom)
    result = tool.run(args(trio(tmp_path), out_dir=str(out)), RunContext())
    assert result.output_paths == []
    assert dict(result.failures)['写 Excel'] == '盘满了'
    assert '盘满了' in result.summary


def test_a_broken_file_is_reported_but_still_exports(tmp_path):
    out = tmp_path / 'out'
    out.mkdir()
    db = library(tmp_path, person('s1', 'Alpha One'))
    dump(db, 's9.json', '{ 这不是 json')
    result = tool.run(args(db, out_dir=str(out)), RunContext())
    assert len(result.output_paths) == 1
    assert [name for name, _ in result.failures] == ['s9.json']


# ---------------- 与在线产出对拍 ----------------
def workbook_dump(path):
    """工作簿的全部内容：页序、页里的表名、每格的值与它指向哪里。"""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    out = []
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for cell in row:
                link = cell.hyperlink
                cells.append((cell.coordinate, cell.value,
                              getattr(link, 'target', None),
                              getattr(link, 'location', None)))
        out.append((ws.title, sorted(ws.tables), cells))
    return out


def test_offline_export_matches_what_the_fetcher_writes(tmp_path, monkeypatch):
    """同样几个人，离线导出与在线抓取必须写出同一个工作簿。

    库里漏存了字段、读回来串了行、组合顺序变了，都只在这里显出来——两条路径各自
    的测试都只跟自己比。用抓取工具那套假 vndb 起这一轮：一次抓取同时产出在线的
    工作簿和本地库，来源相同，剩下的差异只可能来自库的往返。

    真库上也对过一次（s367+s131+s359，库 09-04 抓、工作簿 09-03 抓）：9 张页的
    页序、表名、5 万余格的值与超链接全部逐格相同，差异只有 4 格 Role 和概览页
    跟着变的那几个计数——那是 vndb 上游改了角色主次，不是往返丢了东西。
    """
    pytest.importorskip('openpyxl')
    # 复用 test_vndb_voiced 里的假传输与那个迷你 vndb，别抄第二份数据。
    import test_vndb_voiced as online

    online.FakeApi(online.vndb()).install(monkeypatch)
    db, live, offline = tmp_path / 'db', tmp_path / 'live', tmp_path / 'offline'
    for d in (db, live, offline):
        d.mkdir()

    who = 's1, s2, s3'
    fetched = online.tool.run(online.args(who, live, save_db=True,
                                          db_dir=str(db)), RunContext())
    exported = tool.run(args(str(db), who=who, out_dir=str(offline)),
                        RunContext())
    assert fetched.failures == [] and exported.failures == []
    assert (os.path.basename(fetched.output_paths[0])
            == os.path.basename(exported.output_paths[0]))
    assert workbook_dump(exported.output_paths[0]) == workbook_dump(
        fetched.output_paths[0])


# ---------------- 校验与置灰 ----------------
def test_validate_only_looks_at_paths_and_the_person_cap(tmp_path):
    assert tool.validate(args(str(tmp_path))) == []
    errors = dict(tool.validate(args(str(tmp_path / 'nope'),
                                     out_dir=str(tmp_path / 'nope2'))))
    assert errors['db_dir'] == '目录不存在或不可访问'
    assert errors['out_dir'] == '目录不存在或不可访问'
    # 关键词与人名任意乱填都不算错：库里有没有这个人是 preview 的事。
    assert tool.validate(args(str(tmp_path), who='!!', keyword='!!')) == []


def test_validate_shares_the_person_cap_with_the_fetcher(tmp_path):
    """离线导出也要写 2^N 张工作表，贵的不是抓取。"""
    who = ', '.join('s%d' % i for i in range(1, 11))
    message = dict(tool.validate(args(str(tmp_path), who=who)))['who']
    assert '最多 8 个人' in message


def test_an_empty_export_dir_greys_out_start_but_still_shows_the_view(tmp_path):
    """空着的导出目录不能进 validate：那会连预览一起拦掉，翻库就没法翻了。"""
    db = trio(tmp_path)
    assert tool.validate(args(db)) == []
    blocked = look(db)
    assert not blocked.ok
    assert '只看不导' in blocked.summary
    assert len(blocked.table.rows) == 3
    assert look(db, out_dir=str(tmp_path)).ok


def test_tool_spec_fields_match_what_run_reads(tmp_path):
    keys = [f.key for f in tool.TOOL.fields]
    assert keys == ['db_dir', 'who', 'keyword', 'out_dir']
    # 一个 rescan 字段都没有：读库很便宜，改哪一格都该立刻重算。
    assert [f.key for f in tool.TOOL.fields if f.rescan] == []
    assert [f.key for f in tool.TOOL.fields if f.required] == ['db_dir']
    assert tool.TOOL.category == '资料'
